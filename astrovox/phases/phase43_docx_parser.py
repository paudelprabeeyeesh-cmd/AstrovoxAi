"""DOCX (Microsoft Word) document parser."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class DOCXParagraph:
    """A paragraph in a DOCX document."""
    text: str
    is_heading: bool = False
    level: int = 0

    def to_dict(self) -> dict:
        return {"text": self.text, "is_heading": self.is_heading, "level": self.level}


@dataclass
class DOCXDocument:
    """A parsed DOCX document."""
    path: str
    paragraphs: list[DOCXParagraph] = field(default_factory=list)
    title: str = ""
    metadata: dict = field(default_factory=dict)

    def to_text(self) -> str:
        lines = []
        for p in self.paragraphs:
            if p.is_heading:
                lines.append(f"{'#' * (p.level + 1)} {p.text}")
            else:
                lines.append(p.text)
        return "\n\n".join(lines)

    def to_dict(self) -> dict:
        return {"path": self.path, "title": self.title,
                "paragraphs": [p.to_dict() for p in self.paragraphs[:50]],
                "metadata": self.metadata}


class DOCXParser:
    """Parse DOCX files into structured text."""

    def __init__(self, use_external: bool = False):
        self.use_external = use_external
        self.last_error: str | None = None

    def parse(self, path: str | Path) -> DOCXDocument:
        """Parse a DOCX file."""
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"DOCX not found: {path}")
        try:
            return self._parse_with_lib(path)
        except Exception as e:
            self.last_error = str(e)
            return self._parse_fallback(path)

    def _parse_with_lib(self, path: Path) -> DOCXDocument:
        import docx
        doc = docx.Document(str(path))
        parsed = DOCXDocument(path=str(path))
        for para in doc.paragraphs:
            style = para.style.name if para.style else ""
            parsed.paragraphs.append(DOCXParagraph(
                text=para.text,
                is_heading="Heading" in style,
                level=int(style.split()[-1]) if "Heading" in style else 0))
        parsed.title = doc.core_properties.title or ""
        parsed.metadata = {k: v for k, v in doc.core_properties.__dict__.items() if v}
        return parsed

    def _parse_fallback(self, path: Path) -> DOCXDocument:
        data = path.read_bytes()
        text = data.decode("utf-8", errors="ignore")
        paragraphs = [l for l in text.split("\n") if l.strip()]
        parsed = DOCXDocument(path=str(path), paragraphs=[
            DOCXParagraph(text=p) for p in paragraphs])
        return parsed

    def parse_many(self, paths: list[str]) -> list[DOCXDocument]:
        """Parse multiple DOCX files."""
        return [self.parse(p) for p in paths]

    def extract_text(self, path: str | Path) -> str:
        """Extract all text from a DOCX."""
        return self.parse(path).to_text()
